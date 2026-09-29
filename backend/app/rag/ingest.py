from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import requests
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

_WHITESPACE = re.compile(r"\s+")


@dataclass
class Chunk:
    text: str
    page: int | None = None
    chunk_index: int = 0


@dataclass
class LoadedDocument:
    """A source split into per-section (usually per-page) text before chunking."""

    title: str
    kind: str  # pdf | docx | text | url
    sections: list[tuple[int | None, str]] = field(default_factory=list)


def _clean(text: str) -> str:
    return _WHITESPACE.sub(" ", text).strip()


# ── loaders ──────────────────────────────────────────────────────────────────
def load_pdf(data: bytes, filename: str) -> LoadedDocument:
    reader = PdfReader(io.BytesIO(data))
    sections: list[tuple[int | None, str]] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if text:
            sections.append((page_number, text))
    if not sections:
        raise ValueError("No extractable text found in this PDF (it may be scanned images).")
    return LoadedDocument(title=filename, kind="pdf", sections=sections)


def load_docx(data: bytes, filename: str) -> LoadedDocument:
    doc = DocxDocument(io.BytesIO(data))
    text = _clean("\n".join(p.text for p in doc.paragraphs))
    if not text:
        raise ValueError("No text found in this Word document.")
    return LoadedDocument(title=filename, kind="docx", sections=[(None, text)])


def load_text(data: bytes, filename: str) -> LoadedDocument:
    text = _clean(data.decode("utf-8", errors="replace"))
    if not text:
        raise ValueError("This file is empty.")
    return LoadedDocument(title=filename, kind="text", sections=[(None, text)])


def load_url(url: str, timeout: float) -> LoadedDocument:
    if not url.startswith(("http://", "https://")):
        raise ValueError("URL must start with http:// or https://")
    resp = requests.get(
        url, timeout=timeout, headers={"User-Agent": "Mozilla/5.0 (compatible; AgenticRAG/1.0)"}
    )
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")
    for element in soup(["script", "style", "nav", "footer", "header", "aside", "form"]):
        element.decompose()
    title = _clean(soup.title.string) if soup.title and soup.title.string else url
    text = _clean(" ".join(soup.stripped_strings))
    if not text:
        raise ValueError("Could not extract readable text from that page.")
    return LoadedDocument(title=title, kind="url", sections=[(None, text)])


# ── chunking ─────────────────────────────────────────────────────────────────
def split_text(text: str, chunk_size: int, chunk_overlap: int) -> list[str]:
    if chunk_size <= 0 or chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError("chunk_size must be positive and larger than chunk_overlap")
    text = _clean(text)
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + chunk_size, n)
        if end < n:
            # prefer to break on a sentence/word boundary in the back half
            window = text[start + chunk_size // 2 : end]
            for sep in (". ", "? ", "! ", "\n", " "):
                pos = window.rfind(sep)
                if pos != -1:
                    end = start + chunk_size // 2 + pos + len(sep)
                    break
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= n:
            break
        start = max(start + 1, end - chunk_overlap)
    return chunks


def chunk_document(doc: LoadedDocument, chunk_size: int, chunk_overlap: int) -> list[Chunk]:
    out: list[Chunk] = []
    index = 0
    for page, section_text in doc.sections:
        for piece in split_text(section_text, chunk_size, chunk_overlap):
            out.append(Chunk(text=piece, page=page, chunk_index=index))
            index += 1
    if not out:
        raise ValueError("Document produced no chunks after cleaning.")
    return out
